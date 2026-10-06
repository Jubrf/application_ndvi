"""
Graphiques Excel natifs (openpyxl) : un graphique par parcelle dans l'export temporel.

Onglet « Graphiques » : les graphiques, deux par ligne.
Onglet « Données graphiques » : un bloc de colonnes par parcelle, une ligne par jour
de la période (date, NDVI lissé, NDVI mesuré retenu, NDVI écarté, seuil bas, seuil haut).
Les cellules vides créent des trous dans les courbes (courbe absente hors mesures,
points seulement aux dates de mesure).
"""
import pandas as pd
from openpyxl.chart import Reference, ScatterChart, Series
from openpyxl.chart.layout import Layout, ManualLayout
from openpyxl.chart.marker import Marker
from openpyxl.utils import get_column_letter

MAX_CHARTS = 150
COLS_PER_BLOCK = 6
CURVE_HEX = "2A78D6"
EXCL_HEX = "9E9D98"
THRESH_HEX = "6E6D69"


def _line_series(ref_y, ref_x, title, color, width_emu=22000, dash=None):
    s = Series(ref_y, ref_x, title=title)
    s.marker = Marker(symbol="none")
    s.smooth = False
    s.graphicalProperties.line.solidFill = color
    s.graphicalProperties.line.width = width_emu
    if dash:
        s.graphicalProperties.line.dashStyle = dash
    return s


def _point_series(ref_y, ref_x, title, color, filled=True):
    s = Series(ref_y, ref_x, title=title)
    s.marker = Marker(symbol="circle", size=6)
    if filled:
        s.marker.graphicalProperties.solidFill = color
    else:
        s.marker.graphicalProperties.noFill = True
    s.marker.graphicalProperties.line.solidFill = color
    s.graphicalProperties.line.noFill = True
    return s


def add_parcel_charts(book, synthese, detail, courbes, period_start, period_end, low, high,
                      max_charts=MAX_CHARTS):
    """Ajoute les onglets « Graphiques » et « Données graphiques » au classeur."""
    days = pd.date_range(pd.Timestamp(period_start), pd.Timestamp(period_end), freq="D")
    ws_data = book.create_sheet("Données graphiques")
    ws_chart = book.create_sheet("Graphiques")
    ws_chart["A1"] = ("Un graphique par parcelle : courbe lissée (bleu), mesures retenues "
                      "(points bleus), mesures écartées (cercles gris), seuils (tirets).")
    pids = list(synthese["ID"])[:max_charts]
    if len(synthese) > max_charts:
        ws_chart["A2"] = (f"Limité aux {max_charts} premières parcelles "
                          f"({len(synthese)} au total).")

    n = len(days)
    for k, pid in enumerate(pids):
        c0 = k * COLS_PER_BLOCK + 1
        headers = ["Date", f"{pid} lissé", f"{pid} retenue", f"{pid} écartée", "Seuil bas", "Seuil haut"]
        for j, h in enumerate(headers):
            ws_data.cell(row=1, column=c0 + j, value=h)

        cur = courbes[courbes["ID"] == pid].set_index("Date")["NDVI_lisse"]
        det = detail[detail["ID"] == pid]
        kept = det[det["Retenue"] == "Oui"].set_index("Date")["NDVI"]
        brut = det["NDVI_brut"] if "NDVI_brut" in det else det["NDVI"]
        excl = det.assign(v=brut)[det["Retenue"] != "Oui"].dropna(subset=["v"]).set_index("Date")["v"]

        for i, day in enumerate(days, start=2):
            ws_data.cell(row=i, column=c0, value=day.to_pydatetime()).number_format = "dd/mm/yyyy"
            for j, src in ((1, cur), (2, kept), (3, excl)):
                v = src.get(day)
                if v is not None and not pd.isna(v):
                    ws_data.cell(row=i, column=c0 + j, value=float(v))
            ws_data.cell(row=i, column=c0 + 4, value=float(low))
            ws_data.cell(row=i, column=c0 + 5, value=float(high))
        ws_data.column_dimensions[get_column_letter(c0)].width = 11

        x = Reference(ws_data, min_col=c0, min_row=2, max_row=n + 1)
        ref = lambda j: Reference(ws_data, min_col=c0 + j, min_row=1, max_row=n + 1)

        ch = ScatterChart()
        s_row = synthese[synthese["ID"] == pid].iloc[0]
        phase = s_row.get("Phase_fin")
        ch.title = f"{pid} — fin de période : {phase if isinstance(phase, str) else '—'}"
        ch.style = 2
        ch.height, ch.width = 7.5, 16
        ch.y_axis.title = "NDVI"
        ch.y_axis.scaling.min, ch.y_axis.scaling.max = 0, 1
        ch.y_axis.majorUnit = 0.1
        ch.y_axis.number_format = "0.0"
        ch.x_axis.number_format = "dd/mm/yy"
        ch.x_axis.scaling.min = float((days[0] - pd.Timestamp("1899-12-30")).days)
        ch.x_axis.scaling.max = float((days[-1] - pd.Timestamp("1899-12-30")).days)
        ch.x_axis.majorUnit = 30 if n > 62 else 7
        ch.x_axis.delete = False
        ch.y_axis.delete = False
        ch.display_blanks = "gap"
        ch.y_axis.majorGridlines = None   # les seuils restent les seules lignes horizontales
        ch.x_axis.majorGridlines = None
        ch.legend.position = "b"
        ch.plot_area.layout = Layout(manualLayout=ManualLayout(x=0.08, y=0.12, w=0.88, h=0.62))

        ch.series.append(_line_series(ref(4), x, "Seuil bas", THRESH_HEX, 12700, "dash"))
        ch.series.append(_line_series(ref(5), x, "Seuil haut", THRESH_HEX, 12700, "sysDash"))
        ch.series.append(_line_series(ref(1), x, "NDVI lissé", CURVE_HEX))
        ch.series.append(_point_series(ref(2), x, "Mesure retenue", CURVE_HEX, filled=True))
        ch.series.append(_point_series(ref(3), x, "Mesure écartée", EXCL_HEX, filled=False))

        row = 4 + (k // 2) * 16
        col = 1 if k % 2 == 0 else 11
        ws_chart.add_chart(ch, f"{get_column_letter(col)}{row}")
